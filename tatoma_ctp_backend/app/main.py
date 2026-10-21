import csv
import io
import json
from contextlib import asynccontextmanager
from typing import List, Optional
from fastapi import APIRouter, Depends, FastAPI, File, HTTPException, Query, UploadFile, WebSocket, WebSocketDisconnect
from fastapi.middleware.cors import CORSMiddleware
from fastapi.security import HTTPAuthorizationCredentials
from openpyxl import load_workbook
from sqlalchemy.orm import Session

from app.auth import (
    REVOKED_TOKENS,
    create_jwt_token,
    get_current_user,
    hash_password,
    require_roles,
    security_scheme,
    verify_password,
)
from app.database import Base, SessionLocal, engine, get_db
from app.models import (
    AreaModel,
    CapacityAdjustmentModel,
    OrderAllocationModel,
    OrderModel,
    PlantModel,
    StationModel,
    UserModel,
    WorkerModel,
    WorkflowStepModel,
)
from app.schemas import (
    AreaItem,
    CapacityAdjustmentItem,
    CapacityAdjustmentsPut,
    LoginRequest,
    OrderCreate,
    OrderUpdate,
    PlantLayoutUpdate,
    SimulationApplyRequest,
    SimulationEvaluateResponse,
    StationCreate,
    StationSkillsPut,
    StationStatusPatch,
    StationUpdate,
    TokenResponse,
    UserCreate,
    UserOut,
    UserUpdate,
    WorkerAbsencesUpdate,
    WorkerCreate,
    WorkerOut,
    WorkerUpdate,
    WorkflowStepItem,
)
from app.seed import seed_initial_data
from app.services.ai_local import evaluate_with_local_ai
from app.services.engine import IndustrialPlanningEngine
from app.utils_weeks import format_week, generate_week_range, parse_week
from app.ws_manager import ws_manager


@asynccontextmanager
async def lifespan(app: FastAPI):
    Base.metadata.create_all(bind=engine)
    db = SessionLocal()
    try:
        seed_initial_data(db)
    finally:
        db.close()
    yield


app = FastAPI(
    title="API CTP Industrial Configurable + IA Local (SQLite)",
    version="2.0.0",
    lifespan=lifespan,
)

app.add_middleware(
    CORSMiddleware,
    allow_origins=["*"],
    allow_credentials=True,
    allow_methods=["*"],
    allow_headers=["*"],
)

router = APIRouter(prefix="/api/v1")


# =====================================================================
# 1. AUTENTICACIÓN (/api/v1/auth/* y /api/v1/me)
# =====================================================================
@router.post("/auth/login", response_model=TokenResponse, tags=["Autenticación"])
def login(req: LoginRequest, db: Session = Depends(get_db)):
    user = db.query(UserModel).filter(UserModel.username == req.username).first()
    if not user or not verify_password(req.password, user.password_hash):
        raise HTTPException(status_code=401, detail="Credenciales incorrectas")
    access_token = create_jwt_token({"sub": user.username, "role": user.role})
    refresh_token = create_jwt_token({"sub": user.username, "role": user.role, "type": "refresh"})
    return TokenResponse(
        access_token=access_token,
        refresh_token=refresh_token,
        user=UserOut(id=user.id, username=user.username, full_name=user.full_name, role=user.role),
    )


@router.post("/auth/refresh", response_model=TokenResponse, tags=["Autenticación"])
def refresh_token(current_user: UserModel = Depends(get_current_user)):
    new_access = create_jwt_token({"sub": current_user.username, "role": current_user.role})
    new_refresh = create_jwt_token({"sub": current_user.username, "role": current_user.role, "type": "refresh"})
    return TokenResponse(
        access_token=new_access,
        refresh_token=new_refresh,
        user=UserOut(id=current_user.id, username=current_user.username, full_name=current_user.full_name, role=current_user.role),
    )


@router.post("/auth/logout", tags=["Autenticación"])
def logout(creds: Optional[HTTPAuthorizationCredentials] = Depends(security_scheme)):
    if creds:
        REVOKED_TOKENS.add(creds.credentials)
    return {"status": "logged_out"}


@router.get("/me", response_model=UserOut, tags=["Autenticación"])
@router.get("/auth/me", response_model=UserOut, tags=["Autenticación"])
def get_me(current_user: UserModel = Depends(get_current_user)):
    return UserOut(id=current_user.id, username=current_user.username, full_name=current_user.full_name, role=current_user.role)


# =====================================================================
# 2. USUARIOS Y TRABAJADORES (ADMINISTRADOR)
# =====================================================================
@router.get("/users", response_model=List[UserOut], tags=["Usuarios y Trabajadores"])
def list_users(db: Session = Depends(get_db), _: UserModel = Depends(require_roles("admin"))):
    users = db.query(UserModel).all()
    return [UserOut(id=u.id, username=u.username, full_name=u.full_name, role=u.role) for u in users]


@router.post("/users", response_model=UserOut, tags=["Usuarios y Trabajadores"])
def create_user(payload: UserCreate, db: Session = Depends(get_db), _: UserModel = Depends(require_roles("admin"))):
    if db.query(UserModel).filter(UserModel.username == payload.username).first():
        raise HTTPException(status_code=400, detail="El nombre de usuario ya existe")
    u = UserModel(
        username=payload.username,
        password_hash=hash_password(payload.password),
        full_name=payload.full_name,
        role=payload.role,
    )
    db.add(u)
    db.commit()
    db.refresh(u)
    return UserOut(id=u.id, username=u.username, full_name=u.full_name, role=u.role)


@router.patch("/users/{user_id}", response_model=UserOut, tags=["Usuarios y Trabajadores"])
def update_user(user_id: int, payload: UserUpdate, db: Session = Depends(get_db), _: UserModel = Depends(require_roles("admin"))):
    u = db.query(UserModel).filter(UserModel.id == user_id).first()
    if not u:
        raise HTTPException(status_code=404, detail="Usuario no encontrado")
    if payload.full_name is not None:
        u.full_name = payload.full_name
    if payload.role is not None:
        u.role = payload.role
    if payload.password:
        u.password_hash = hash_password(payload.password)
    db.commit()
    db.refresh(u)
    return UserOut(id=u.id, username=u.username, full_name=u.full_name, role=u.role)


@router.delete("/users/{user_id}", tags=["Usuarios y Trabajadores"])
def delete_user(user_id: int, db: Session = Depends(get_db), _: UserModel = Depends(require_roles("admin"))):
    u = db.query(UserModel).filter(UserModel.id == user_id).first()
    if not u:
        raise HTTPException(status_code=404, detail="Usuario no encontrado")
    db.delete(u)
    db.commit()
    return {"status": "deleted", "id": user_id}


def _serialize_worker(w: WorkerModel) -> WorkerOut:
    return WorkerOut(
        id=w.id,
        name=w.name,
        station_id=w.station_id,
        skills=json.loads(w.skills_json or '["general"]'),
        absences=json.loads(w.absences_json or "[]"),
    )


@router.get("/workers", response_model=List[WorkerOut], tags=["Usuarios y Trabajadores"])
def list_workers(db: Session = Depends(get_db), _: UserModel = Depends(get_current_user)):
    return [_serialize_worker(w) for w in db.query(WorkerModel).all()]


@router.post("/workers", response_model=WorkerOut, tags=["Usuarios y Trabajadores"])
def create_worker(payload: WorkerCreate, db: Session = Depends(get_db), _: UserModel = Depends(require_roles("admin"))):
    w = WorkerModel(
        name=payload.name,
        station_id=payload.station_id,
        skills_json=json.dumps(payload.skills),
        absences_json="[]",
    )
    db.add(w)
    db.commit()
    db.refresh(w)
    return _serialize_worker(w)


@router.patch("/workers/{worker_id}", response_model=WorkerOut, tags=["Usuarios y Trabajadores"])
def update_worker(worker_id: int, payload: WorkerUpdate, db: Session = Depends(get_db), _: UserModel = Depends(require_roles("admin"))):
    w = db.query(WorkerModel).filter(WorkerModel.id == worker_id).first()
    if not w:
        raise HTTPException(status_code=404, detail="Trabajador no encontrado")
    if payload.name is not None:
        w.name = payload.name
    if payload.station_id is not None:
        w.station_id = payload.station_id
    if payload.skills is not None:
        w.skills_json = json.dumps(payload.skills)
    db.commit()
    db.refresh(w)
    return _serialize_worker(w)


@router.delete("/workers/{worker_id}", tags=["Usuarios y Trabajadores"])
def delete_worker(worker_id: int, db: Session = Depends(get_db), _: UserModel = Depends(require_roles("admin"))):
    w = db.query(WorkerModel).filter(WorkerModel.id == worker_id).first()
    if not w:
        raise HTTPException(status_code=404, detail="Trabajador no encontrado")
    db.delete(w)
    db.commit()
    return {"status": "deleted", "id": worker_id}


@router.put("/workers/{worker_id}/absences", response_model=WorkerOut, tags=["Usuarios y Trabajadores"])
def put_worker_absences(
    worker_id: int,
    payload: WorkerAbsencesUpdate,
    db: Session = Depends(get_db),
    _: UserModel = Depends(require_roles("admin", "planner")),
):
    w = db.query(WorkerModel).filter(WorkerModel.id == worker_id).first()
    if not w:
        raise HTTPException(status_code=404, detail="Trabajador no encontrado")
    w.absences_json = json.dumps([a.model_dump() for a in payload.absences])
    db.commit()
    db.refresh(w)
    return _serialize_worker(w)


# =====================================================================
# 3. PLANTA, PUESTOS, ÁREAS Y FLUJO DE TRABAJO CONFIGURABLE
# =====================================================================
def _serialize_station(s: StationModel) -> dict:
    return {
        "id": s.id,
        "plant_id": s.plant_id,
        "name": s.name,
        "area": s.area,
        "x": s.x,
        "y": s.y,
        "w": s.w,
        "h": s.h,
        "workers": s.workers,
        "shifts": s.shifts,
        "max_shifts": s.max_shifts,
        "hoursPerShift": s.hours_per_shift,
        "unitsPerShift": s.units_per_shift,
        "status": s.status,
        "skills": json.loads(s.skills_json or '["general"]'),
    }


@router.get("/plants/{plant_id}", tags=["Planta y Puestos"])
def get_plant(plant_id: str, db: Session = Depends(get_db), _: UserModel = Depends(get_current_user)):
    plant = db.query(PlantModel).filter(PlantModel.id == plant_id).first()
    if not plant:
        plant = db.query(PlantModel).first()
    stations = db.query(StationModel).filter(StationModel.plant_id == plant.id).all()
    return {
        "id": plant.id,
        "name": plant.name,
        "cols": plant.cols,
        "rows": plant.rows,
        "stations": [_serialize_station(s) for s in stations],
    }


@router.put("/plants/{plant_id}/layout", tags=["Planta y Puestos"])
def update_plant_layout(
    plant_id: str,
    payload: PlantLayoutUpdate,
    db: Session = Depends(get_db),
    _: UserModel = Depends(require_roles("admin")),
):
    plant = db.query(PlantModel).filter(PlantModel.id == plant_id).first()
    if not plant:
        plant = PlantModel(id=plant_id, name=payload.name or plant_id, cols=payload.cols, rows=payload.rows)
        db.add(plant)
    else:
        plant.cols = payload.cols
        plant.rows = payload.rows
        if payload.name:
            plant.name = payload.name

    if payload.positions:
        for pos in payload.positions:
            st = db.query(StationModel).filter(StationModel.id == pos.id).first()
            if st:
                st.x, st.y, st.w, st.h = pos.x, pos.y, pos.w, pos.h

    db.commit()
    return get_plant(plant.id, db)


@router.get("/stations", tags=["Planta y Puestos"])
def list_stations(db: Session = Depends(get_db), _: UserModel = Depends(get_current_user)):
    return [_serialize_station(s) for s in db.query(StationModel).all()]


@router.post("/stations", tags=["Planta y Puestos"])
async def create_station(
    payload: StationCreate,
    db: Session = Depends(get_db),
    _: UserModel = Depends(require_roles("admin")),
):
    if db.query(StationModel).filter(StationModel.id == payload.id).first():
        raise HTTPException(status_code=400, detail="El ID de puesto ya existe")
    if not db.query(AreaModel).filter(AreaModel.id == payload.area).first():
        db.add(AreaModel(id=payload.area, name=payload.area.capitalize()))

    st = StationModel(
        id=payload.id,
        name=payload.name,
        area=payload.area,
        x=payload.x,
        y=payload.y,
        w=payload.w,
        h=payload.h,
        workers=payload.workers,
        shifts=payload.shifts,
        max_shifts=payload.max_shifts,
        hours_per_shift=payload.hoursPerShift,
        units_per_shift=payload.unitsPerShift,
        status=payload.status,
        skills_json=json.dumps(payload.skills),
    )
    db.add(st)
    db.commit()
    db.refresh(st)
    res = _serialize_station(st)
    await ws_manager.broadcast("station_created", res)
    return res


@router.patch("/stations/{station_id}", tags=["Planta y Puestos"])
async def update_station(
    station_id: str,
    payload: StationUpdate,
    db: Session = Depends(get_db),
    _: UserModel = Depends(require_roles("admin")),
):
    st = db.query(StationModel).filter(StationModel.id == station_id).first()
    if not st:
        raise HTTPException(status_code=404, detail="Puesto no encontrado")

    for field, attr in [
        ("name", "name"),
        ("area", "area"),
        ("x", "x"),
        ("y", "y"),
        ("w", "w"),
        ("h", "h"),
        ("workers", "workers"),
        ("shifts", "shifts"),
        ("max_shifts", "max_shifts"),
        ("hoursPerShift", "hours_per_shift"),
        ("unitsPerShift", "units_per_shift"),
        ("status", "status"),
    ]:
        val = getattr(payload, field)
        if val is not None:
            setattr(st, attr, val)

    if payload.skills is not None:
        st.skills_json = json.dumps(payload.skills)

    db.commit()
    db.refresh(st)
    res = _serialize_station(st)
    await ws_manager.broadcast("station_updated", res)
    return res


@router.delete("/stations/{station_id}", tags=["Planta y Puestos"])
async def delete_station(
    station_id: str,
    db: Session = Depends(get_db),
    _: UserModel = Depends(require_roles("admin")),
):
    st = db.query(StationModel).filter(StationModel.id == station_id).first()
    if not st:
        raise HTTPException(status_code=404, detail="Puesto no encontrado")
    db.delete(st)
    db.commit()
    await ws_manager.broadcast("station_deleted", {"id": station_id})
    return {"status": "deleted", "id": station_id}


@router.put("/stations/{station_id}/skills", tags=["Planta y Puestos"])
def update_station_skills(
    station_id: str,
    payload: StationSkillsPut,
    db: Session = Depends(get_db),
    _: UserModel = Depends(require_roles("admin")),
):
    st = db.query(StationModel).filter(StationModel.id == station_id).first()
    if not st:
        raise HTTPException(status_code=404, detail="Puesto no encontrado")
    st.skills_json = json.dumps(payload.skills)
    db.commit()
    db.refresh(st)
    return _serialize_station(st)


@router.patch("/stations/{station_id}/status", tags=["Capacidad y Disponibilidad"])
async def patch_station_status(
    station_id: str,
    payload: StationStatusPatch,
    db: Session = Depends(get_db),
    _: UserModel = Depends(require_roles("admin", "planner")),
):
    st = db.query(StationModel).filter(StationModel.id == station_id).first()
    if not st:
        raise HTTPException(status_code=404, detail="Puesto no encontrado")
    st.status = payload.status
    db.commit()
    db.refresh(st)
    res = _serialize_station(st)
    await ws_manager.broadcast("station_status_changed", res)
    return res


@router.get("/areas", response_model=List[AreaItem], tags=["Planta y Puestos"])
def get_areas(db: Session = Depends(get_db), _: UserModel = Depends(get_current_user)):
    areas = db.query(AreaModel).all()
    return [
        AreaItem(
            id=a.id,
            name=a.name,
            color=a.color,
            hourly_cost=a.hourly_cost,
            overtime_hourly_cost=a.overtime_hourly_cost,
            subcontract_hourly_cost=a.subcontract_hourly_cost,
            subcontract_fixed_logistics=a.subcontract_fixed_logistics,
            subcontract_lead_time_weeks=a.subcontract_lead_time_weeks,
        )
        for a in areas
    ]


@router.put("/areas", response_model=List[AreaItem], tags=["Planta y Puestos"])
def put_areas(
    areas: List[AreaItem],
    db: Session = Depends(get_db),
    _: UserModel = Depends(require_roles("admin")),
):
    db.query(AreaModel).delete()
    for a in areas:
        db.add(AreaModel(**a.model_dump()))
    db.commit()
    return get_areas(db)


@router.get("/workflow", response_model=List[WorkflowStepItem], tags=["Planta y Puestos"])
def get_workflow(db: Session = Depends(get_db), _: UserModel = Depends(get_current_user)):
    steps = db.query(WorkflowStepModel).order_by(WorkflowStepModel.step_order.asc()).all()
    return [
        WorkflowStepItem(
            step_order=s.step_order,
            area_id=s.area_id,
            name=s.name,
            offset_weeks_before_delivery=s.offset_weeks_before_delivery,
            is_optional=s.is_optional,
        )
        for s in steps
    ]


@router.put("/workflow", response_model=List[WorkflowStepItem], tags=["Planta y Puestos"])
def put_workflow(
    steps: List[WorkflowStepItem],
    db: Session = Depends(get_db),
    _: UserModel = Depends(require_roles("admin")),
):
    db.query(WorkflowStepModel).delete()
    for s in sorted(steps, key=lambda x: x.step_order):
        if not db.query(AreaModel).filter(AreaModel.id == s.area_id).first():
            db.add(AreaModel(id=s.area_id, name=s.name))
        db.add(WorkflowStepModel(**s.model_dump()))
    db.commit()
    return get_workflow(db)


# =====================================================================
# 4. CAPACIDAD Y DISPONIBILIDAD
# =====================================================================
@router.get("/capacity", tags=["Capacidad y Disponibilidad"])
def get_capacity(
    from_week: str = Query("2026-W01", alias="from"),
    to_week: str = Query("2026-W36", alias="to"),
    db: Session = Depends(get_db),
    _: UserModel = Depends(get_current_user),
):
    eng = IndustrialPlanningEngine(db)
    weeks = generate_week_range(from_week, to_week)
    result = []
    for w in weeks:
        by_area = {}
        stations_out = []
        for st in eng.stations:
            net_h, max_u = eng.get_station_net_capacity(st, w)
            comm_h, comm_u = eng.get_station_committed_load(st.id, w)
            by_area.setdefault(st.area, {"capacity_hours": 0.0, "committed_hours": 0.0})
            by_area[st.area]["capacity_hours"] = round(by_area[st.area]["capacity_hours"] + net_h, 2)
            by_area[st.area]["committed_hours"] = round(by_area[st.area]["committed_hours"] + comm_h, 2)
            stations_out.append(
                {
                    "station_id": st.id,
                    "area": st.area,
                    "capacity_hours": net_h,
                    "committed_hours": comm_h,
                    "max_units": max_u,
                    "committed_units": comm_u,
                }
            )
        result.append({"week": w, "areas": by_area, "stations": stations_out})
    return result


@router.get("/capacity/adjustments", response_model=List[CapacityAdjustmentItem], tags=["Capacidad y Disponibilidad"])
def get_capacity_adjustments(db: Session = Depends(get_db), _: UserModel = Depends(get_current_user)):
    adjs = db.query(CapacityAdjustmentModel).all()
    return [
        CapacityAdjustmentItem(
            id=a.id, week=a.week, area=a.area, station_id=a.station_id, hours_delta=a.hours_delta, reason=a.reason
        )
        for a in adjs
    ]


@router.put("/capacity/adjustments", response_model=List[CapacityAdjustmentItem], tags=["Capacidad y Disponibilidad"])
async def put_capacity_adjustments(
    payload: CapacityAdjustmentsPut,
    db: Session = Depends(get_db),
    _: UserModel = Depends(require_roles("admin", "planner")),
):
    db.query(CapacityAdjustmentModel).delete()
    for a in payload.adjustments:
        y, wn = parse_week(a.week)
        db.add(
            CapacityAdjustmentModel(
                week=format_week(y, wn),
                area=a.area,
                station_id=a.station_id,
                hours_delta=a.hours_delta,
                reason=a.reason,
            )
        )
    db.commit()
    res = get_capacity_adjustments(db)
    await ws_manager.broadcast("capacity_adjusted", {"count": len(res)})
    return res


# =====================================================================
# 5. PEDIDOS + IMPORTACIÓN DIRECTA DE EXCEL
# =====================================================================
def _serialize_order(o: OrderModel, db: Session) -> dict:
    allocs = db.query(OrderAllocationModel).filter(OrderAllocationModel.order_id == o.id).all()
    return {
        "id": o.id,
        "code": o.code,
        "client": o.client,
        "quantity": o.quantity,
        "current_week": o.current_week,
        "delivery_week": o.delivery_week,
        "lead_time_weeks": o.lead_time_weeks,
        "required_skill": o.required_skill,
        "hours_per_unit": json.loads(o.hours_per_unit_json),
        "active_phases": json.loads(o.active_phases_json),
        "status": o.status,
        "allocations": [
            {
                "week": a.week,
                "area": a.area,
                "station_id": a.station_id,
                "allocated_hours": a.allocated_hours,
                "allocated_units": a.allocated_units,
                "is_subcontracted": a.is_subcontracted,
            }
            for a in allocs
        ],
    }


def _allocate_order_in_db(o: OrderModel, db: Session):
    db.query(OrderAllocationModel).filter(OrderAllocationModel.order_id == o.id).delete()
    eng = IndustrialPlanningEngine(db)
    order_schema = OrderCreate(
        code=o.code,
        client=o.client,
        quantity=o.quantity,
        current_week=o.current_week,
        delivery_week=o.delivery_week,
        lead_time_weeks=o.lead_time_weeks,
        required_skill=o.required_skill,
        hours_per_unit=json.loads(o.hours_per_unit_json),
        active_phases=json.loads(o.active_phases_json),
        status=o.status,
    )
    _, _, _, baseline_allocs = eng.evaluate_baseline(order_schema)
    if baseline_allocs:
        for al in baseline_allocs:
            db.add(
                OrderAllocationModel(
                    order_id=o.id,
                    order_code=o.code,
                    client=o.client,
                    week=al.week,
                    area=al.area,
                    station_id=al.station_id,
                    allocated_hours=al.allocated_hours,
                    allocated_units=al.allocated_units,
                    is_subcontracted=al.is_subcontracted,
                )
            )
    else:
        # Si el pedido se guarda forzado, asigna a la primera línea del área para reflejar el exceso en rojo en el plano
        for step in eng.get_active_workflow_steps(order_schema):
            pw = format_week(*parse_week(o.delivery_week))
            from app.utils_weeks import shift_week
            pw = shift_week(o.delivery_week, -step.offset_weeks_before_delivery)
            compat = eng.get_compatible_stations(step.area_id, o.required_skill)
            st_id = compat[0].id if compat else step.area_id.upper()
            hpu = order_schema.hours_per_unit.get(step.area_id, 0.0)
            db.add(
                OrderAllocationModel(
                    order_id=o.id,
                    order_code=o.code,
                    client=o.client,
                    week=pw,
                    area=step.area_id,
                    station_id=st_id,
                    allocated_hours=round(o.quantity * hpu, 2),
                    allocated_units=o.quantity,
                    is_subcontracted=False,
                )
            )


@router.get("/orders", tags=["Pedidos"])
def list_orders(db: Session = Depends(get_db), _: UserModel = Depends(get_current_user)):
    return [_serialize_order(o, db) for o in db.query(OrderModel).all()]


@router.get("/orders/{order_id}", tags=["Pedidos"])
def get_order(order_id: int, db: Session = Depends(get_db), _: UserModel = Depends(get_current_user)):
    o = db.query(OrderModel).filter(OrderModel.id == order_id).first()
    if not o:
        raise HTTPException(status_code=404, detail="Pedido no encontrado")
    return _serialize_order(o, db)


@router.post("/orders", tags=["Pedidos"])
async def create_order(
    payload: OrderCreate,
    db: Session = Depends(get_db),
    _: UserModel = Depends(require_roles("admin", "planner")),
):
    if db.query(OrderModel).filter(OrderModel.code == payload.code).first():
        raise HTTPException(status_code=400, detail=f"El pedido '{payload.code}' ya existe")
    active_phases = payload.active_phases or [k for k, v in payload.hours_per_unit.items() if v > 0]
    y, wn = parse_week(payload.delivery_week)
    cy, cwn = parse_week(payload.current_week)

    o = OrderModel(
        code=payload.code,
        client=payload.client,
        quantity=payload.quantity,
        current_week=format_week(cy, cwn),
        delivery_week=format_week(y, wn),
        lead_time_weeks=payload.lead_time_weeks,
        required_skill=payload.required_skill,
        hours_per_unit_json=json.dumps(payload.hours_per_unit),
        active_phases_json=json.dumps(active_phases),
        status=payload.status,
    )
    db.add(o)
    db.flush()
    if o.status == "confirmed":
        _allocate_order_in_db(o, db)
    db.commit()
    db.refresh(o)
    res = _serialize_order(o, db)
    await ws_manager.broadcast("order_created", res)
    return res


@router.patch("/orders/{order_id}", tags=["Pedidos"])
async def update_order(
    order_id: int,
    payload: OrderUpdate,
    db: Session = Depends(get_db),
    _: UserModel = Depends(require_roles("admin", "planner")),
):
    o = db.query(OrderModel).filter(OrderModel.id == order_id).first()
    if not o:
        raise HTTPException(status_code=404, detail="Pedido no encontrado")
    if payload.client is not None:
        o.client = payload.client
    if payload.quantity is not None:
        o.quantity = payload.quantity
    if payload.delivery_week is not None:
        o.delivery_week = format_week(*parse_week(payload.delivery_week))
    if payload.lead_time_weeks is not None:
        o.lead_time_weeks = payload.lead_time_weeks
    if payload.required_skill is not None:
        o.required_skill = payload.required_skill
    if payload.hours_per_unit is not None:
        o.hours_per_unit_json = json.dumps(payload.hours_per_unit)
    if payload.active_phases is not None:
        o.active_phases_json = json.dumps(payload.active_phases)
    if payload.status is not None:
        o.status = payload.status

    if o.status == "confirmed":
        _allocate_order_in_db(o, db)
    db.commit()
    db.refresh(o)
    res = _serialize_order(o, db)
    await ws_manager.broadcast("order_updated", res)
    return res


@router.delete("/orders/{order_id}", tags=["Pedidos"])
async def delete_order(
    order_id: int,
    db: Session = Depends(get_db),
    _: UserModel = Depends(require_roles("admin", "planner")),
):
    o = db.query(OrderModel).filter(OrderModel.id == order_id).first()
    if not o:
        raise HTTPException(status_code=404, detail="Pedido no encontrado")
    db.query(OrderAllocationModel).filter(OrderAllocationModel.order_id == o.id).delete()
    db.delete(o)
    db.commit()
    await ws_manager.broadcast("order_deleted", {"id": order_id})
    return {"status": "deleted", "id": order_id}


@router.post("/orders/{order_id}/confirm", tags=["Pedidos"])
async def confirm_order(
    order_id: int,
    db: Session = Depends(get_db),
    _: UserModel = Depends(require_roles("admin", "planner")),
):
    o = db.query(OrderModel).filter(OrderModel.id == order_id).first()
    if not o:
        raise HTTPException(status_code=404, detail="Pedido no encontrado")
    o.status = "confirmed"
    _allocate_order_in_db(o, db)
    db.commit()
    db.refresh(o)
    res = _serialize_order(o, db)
    await ws_manager.broadcast("order_confirmed", res)
    return res


@router.post("/orders/import", tags=["Pedidos"])
async def import_orders_excel(
    file: UploadFile = File(...),
    db: Session = Depends(get_db),
    _: UserModel = Depends(require_roles("admin", "planner")),
):
    """
    Importa pedidos desde un archivo Excel (.xlsx) o CSV.
    Cabeceras soportadas: code, client, quantity, delivery_week, lead_time_weeks, required_skill, soldadura, pintura, montaje
    """
    content = await file.read()
    rows_data = []

    if file.filename and file.filename.lower().endswith(".csv"):
        decoded = content.decode("utf-8-sig")
        reader = csv.DictReader(io.StringIO(decoded))
        rows_data = list(reader)
    else:
        wb = load_workbook(filename=io.BytesIO(content), data_only=True)
        ws = wb.active
        headers = [str(cell.value or "").strip().lower() for cell in next(ws.iter_rows(min_row=1, max_row=1))]
        for row in ws.iter_rows(min_row=2, values_only=True):
            if any(row):
                rows_data.append(dict(zip(headers, row)))

    imported = []
    errors = []
    known_areas = [a.id for a in db.query(AreaModel).all()]

    for idx, r in enumerate(rows_data, start=2):
        try:
            code = str(r.get("code") or r.get("codigo") or f"IMP-{idx}").strip()
            client = str(r.get("client") or r.get("cliente") or "Cliente Excel").strip()
            qty = int(float(r.get("quantity") or r.get("cantidad") or 1))
            deliv = format_week(*parse_week(str(r.get("delivery_week") or r.get("entrega") or "2026-W12")))
            lt = int(float(r.get("lead_time_weeks") or r.get("plazo_previo") or 2))
            skill = str(r.get("required_skill") or r.get("skill") or "general").strip()

            hpu = {}
            active_phases = []
            for area_id in known_areas:
                val = r.get(area_id)
                if val is not None and str(val).strip() != "":
                    fval = float(val)
                    hpu[area_id] = fval
                    if fval > 0:
                        active_phases.append(area_id)

            if not hpu:
                hpu = {"soldadura": 10.0, "pintura": 4.0, "montaje": 6.0}
                active_phases = ["soldadura", "pintura", "montaje"]

            existing = db.query(OrderModel).filter(OrderModel.code == code).first()
            if existing:
                existing.client = client
                existing.quantity = qty
                existing.delivery_week = deliv
                existing.lead_time_weeks = lt
                existing.required_skill = skill
                existing.hours_per_unit_json = json.dumps(hpu)
                existing.active_phases_json = json.dumps(active_phases)
                _allocate_order_in_db(existing, db)
                imported.append(code)
            else:
                new_o = OrderModel(
                    code=code,
                    client=client,
                    quantity=qty,
                    current_week="2026-W01",
                    delivery_week=deliv,
                    lead_time_weeks=lt,
                    required_skill=skill,
                    hours_per_unit_json=json.dumps(hpu),
                    active_phases_json=json.dumps(active_phases),
                    status="confirmed",
                )
                db.add(new_o)
                db.flush()
                _allocate_order_in_db(new_o, db)
                imported.append(code)
        except Exception as e:
            errors.append({"row": idx, "error": str(e)})

    db.commit()
    await ws_manager.broadcast("orders_imported", {"imported_count": len(imported)})
    return {"imported_count": len(imported), "imported_codes": imported, "errors": errors}


# =====================================================================
# 6. MOTOR DE CÁLCULO Y SIMULACIONES + IA LOCAL
# =====================================================================
@router.post("/simulations/evaluate", response_model=SimulationEvaluateResponse, tags=["Motor de Cálculo"])
async def evaluate_simulation(
    order: OrderCreate,
    db: Session = Depends(get_db),
    _: UserModel = Depends(get_current_user),
):
    order.delivery_week = format_week(*parse_week(order.delivery_week))
    order.current_week = format_week(*parse_week(order.current_week))

    eng = IndustrialPlanningEngine(db)
    lead_ok, earliest_w, conflicts, baseline_allocs = eng.evaluate_baseline(order)
    viable = lead_ok and len(conflicts) == 0

    suggestions = eng.generate_5_strategies(order, conflicts, earliest_w)
    ai_verdict = await evaluate_with_local_ai(order, conflicts, suggestions)

    return SimulationEvaluateResponse(
        viable=viable,
        lead_time_ok=lead_ok,
        earliest_allowed_start_week=earliest_w,
        conflicts=conflicts,
        baseline_allocations=baseline_allocs,
        suggestions=suggestions,
        ai_advisor=ai_verdict,
    )


@router.post("/simulations/apply", tags=["Motor de Cálculo"])
async def apply_simulation(
    payload: SimulationApplyRequest,
    db: Session = Depends(get_db),
    _: UserModel = Depends(require_roles("admin", "planner")),
):
    # 1. Crear ajustes de capacidad si la estrategia elegida incluye horas extra / turnos adicionales
    for adj in payload.capacity_adjustments_to_create:
        db.add(
            CapacityAdjustmentModel(
                week=format_week(*parse_week(adj.week)),
                area=adj.area,
                station_id=adj.station_id,
                hours_delta=adj.hours_delta,
                reason=adj.reason,
            )
        )

    # 2. Guardar o actualizar el pedido confirmado con las asignaciones exactas de la estrategia
    order_data = payload.order
    final_deliv = format_week(*parse_week(payload.final_delivery_week or order_data.delivery_week))
    active_phases = order_data.active_phases or [k for k, v in order_data.hours_per_unit.items() if v > 0]

    o = db.query(OrderModel).filter(OrderModel.code == order_data.code).first()
    if not o:
        o = OrderModel(
            code=order_data.code,
            client=order_data.client,
            quantity=order_data.quantity,
            current_week=format_week(*parse_week(order_data.current_week)),
            delivery_week=final_deliv,
            lead_time_weeks=order_data.lead_time_weeks,
            required_skill=order_data.required_skill,
            hours_per_unit_json=json.dumps(order_data.hours_per_unit),
            active_phases_json=json.dumps(active_phases),
            status="confirmed",
        )
        db.add(o)
        db.flush()
    else:
        o.delivery_week = final_deliv
        o.status = "confirmed"
        db.query(OrderAllocationModel).filter(OrderAllocationModel.order_id == o.id).delete()

    for al in payload.allocations:
        db.add(
            OrderAllocationModel(
                order_id=o.id,
                order_code=o.code,
                client=o.client,
                week=format_week(*parse_week(al.week)),
                area=al.area,
                station_id=al.station_id,
                allocated_hours=al.allocated_hours,
                allocated_units=al.allocated_units,
                is_subcontracted=al.is_subcontracted,
            )
        )

    db.commit()
    db.refresh(o)
    res = _serialize_order(o, db)
    await ws_manager.broadcast("simulation_applied", {"order": res, "strategy": payload.selected_strategy_code})
    return {"status": "applied", "selected_strategy": payload.selected_strategy_code, "order": res}


# =====================================================================
# 7. VISUALIZACIÓN DEL PLANO Y TIRA SEMANAL (HEATMAP) + WEBSOCKET
# =====================================================================
@router.get("/plant-status", tags=["Visualización"])
def get_plant_status(
    week: str = Query("2026-W08"),
    db: Session = Depends(get_db),
    _: UserModel = Depends(get_current_user),
):
    target_w = format_week(*parse_week(week))
    eng = IndustrialPlanningEngine(db)

    stations_status = []
    max_load_in_week = 80.0

    for st in eng.stations:
        net_h, max_u = eng.get_station_net_capacity(st, target_w)
        comm_h, comm_u = eng.get_station_committed_load(st.id, target_w)
        max_load_in_week = max(max_load_in_week, net_h, comm_h)

        # Segmentos divididos por pedido para pintar la barra en perspectiva
        segments = []
        for alloc in eng.allocations:
            if not alloc.is_subcontracted and alloc.station_id == st.id and parse_week(alloc.week) == parse_week(target_w):
                segments.append(
                    {
                        "order_id": alloc.order_id,
                        "order_code": alloc.order_code,
                        "client": alloc.client,
                        "hours": alloc.allocated_hours,
                        "units": alloc.allocated_units,
                    }
                )

        hrs_pct = (comm_h / net_h * 100.0) if net_h > 0 else (100.0 if comm_h > 0 else 0.0)
        unit_pct = (comm_u / max_u * 100.0) if (max_u and max_u > 0) else 0.0
        occ_pct = round(max(hrs_pct, unit_pct), 1)

        if occ_pct < 80.0:
            border_color = "green"
        elif occ_pct <= 100.0:
            border_color = "amber"
        else:
            border_color = "red"

        assigned_workers_list = [
            {"id": w.id, "name": w.name} for w in eng.workers if w.station_id == st.id
        ]

        stations_status.append(
            {
                **_serialize_station(st),
                "week": target_w,
                "load_hours": comm_h,
                "capacity_hours": net_h,
                "load_units": comm_u,
                "capacity_units": max_u,
                "occupancy_pct": occ_pct,
                "border_state": border_color,
                "assigned_workers": assigned_workers_list,
                "segments": segments,
            }
        )

    return {
        "week": target_w,
        "bar_scale_max_hours": max_load_in_week,
        "stations": stations_status,
    }


@router.get("/heatmap", tags=["Visualización"])
def get_weekly_heatmap(
    from_week: str = Query("2026-W01", alias="from"),
    to_week: str = Query("2026-W36", alias="to"),
    db: Session = Depends(get_db),
    _: UserModel = Depends(get_current_user),
):
    eng = IndustrialPlanningEngine(db)
    weeks = generate_week_range(from_week, to_week)
    strip = []

    for w in weeks:
        total_cap = 0.0
        total_load = 0.0
        max_station_pct = 0.0

        for st in eng.stations:
            net_h, max_u = eng.get_station_net_capacity(st, w)
            comm_h, comm_u = eng.get_station_committed_load(st.id, w)
            total_cap += net_h
            total_load += comm_h
            st_pct = (comm_h / net_h * 100.0) if net_h > 0 else (100.0 if comm_h > 0 else 0.0)
            if max_u and max_u > 0:
                st_pct = max(st_pct, (comm_u / max_u * 100.0))
            max_station_pct = max(max_station_pct, st_pct)

        overall_pct = round((total_load / total_cap * 100.0), 1) if total_cap > 0 else 0.0
        effective_pct = round(max(overall_pct, max_station_pct), 1)

        if effective_pct < 80.0:
            color = "green"
        elif effective_pct <= 100.0:
            color = "amber"
        else:
            color = "red"

        strip.append(
            {
                "week": w,
                "total_load_hours": round(total_load, 2),
                "total_capacity_hours": round(total_cap, 2),
                "average_occupancy_pct": overall_pct,
                "peak_station_occupancy_pct": effective_pct,
                "status_color": color,
            }
        )

    return strip


@app.websocket("/api/v1/ws/plant")
@app.websocket("/ws/plant")
async def websocket_plant_endpoint(websocket: WebSocket):
    await ws_manager.connect(websocket)
    try:
        while True:
            data = await websocket.receive_text()
            await websocket.send_json({"event": "pong", "received": data})
    except WebSocketDisconnect:
        ws_manager.disconnect(websocket)


app.include_router(router)
