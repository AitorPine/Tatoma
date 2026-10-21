import re
import unicodedata
from datetime import date
from io import BytesIO
from typing import Any, Optional

from openpyxl import load_workbook


def _normalize_header(value: Any) -> str:
    text = unicodedata.normalize("NFKD", str(value or "").strip().lower())
    return "".join(char for char in text if not unicodedata.combining(char))


def _scheduled_quantity(value: Any) -> int:
    if isinstance(value, (int, float)):
        quantity = int(value)
        if quantity > 0 and quantity == value:
            return quantity
        raise ValueError(f"Cantidad semanal no válida: {value}")

    token = re.sub(r"\s+", "", str(value or "").upper())
    special_quantities = {"3D": 1, "2PRO": 2, "1M": 1, "PRO": 1, "PROM": 1}
    if token in special_quantities:
        return special_quantities[token]
    if re.fullmatch(r"\d+", token):
        quantity = int(token)
        if quantity > 0:
            return quantity
    raise ValueError(f"Marcador semanal no reconocido: {value}")


def parse_commercial_workbook(content: bytes) -> tuple[Optional[list[dict[str, Any]]], list[dict[str, Any]]]:
    workbook = load_workbook(filename=BytesIO(content), read_only=True, data_only=True)
    required_headers = {
        "of",
        "fecha",
        "cliente",
        "proyecto",
        "restante",
        "linea produccion",
        "horas soldadura",
        "horas pintura",
        "horas montaje",
    }

    try:
        for sheet in workbook.worksheets:
            header_row = None
            header_columns: dict[str, int] = {}
            for row_number, row in enumerate(sheet.iter_rows(min_row=1, max_row=min(sheet.max_row, 20), values_only=True), start=1):
                normalized = {_normalize_header(value): index for index, value in enumerate(row)}
                if required_headers.issubset(normalized):
                    header_row = row_number
                    header_columns = normalized
                    break

            if header_row is None:
                continue

            header_values = next(
                sheet.iter_rows(min_row=header_row, max_row=header_row, values_only=True)
            )
            week_columns = []
            first_week_one_column = None
            for column_index, value in enumerate(header_values):
                if isinstance(value, (int, float)) and int(value) == value and 1 <= int(value) <= 52:
                    week_columns.append((column_index, int(value)))
                    if int(value) == 1:
                        first_week_one_column = column_index

            if first_week_one_column is None:
                raise ValueError("No se encontró el inicio de las semanas en la cabecera del Excel")

            orders: list[dict[str, Any]] = []
            errors: list[dict[str, Any]] = []
            for row_number, row in enumerate(sheet.iter_rows(min_row=header_row + 1, values_only=True), start=header_row + 1):
                source_of = str(row[header_columns["of"]] or "").strip()
                client = str(row[header_columns["cliente"]] or "").strip()
                project = str(row[header_columns["proyecto"]] or "").strip()
                if not source_of or not client:
                    continue

                hours_per_unit = {
                    "soldadura": float(row[header_columns["horas soldadura"]] or 0),
                    "pintura": float(row[header_columns["horas pintura"]] or 0),
                    "montaje": float(row[header_columns["horas montaje"]] or 0),
                }
                active_phases = [area for area, hours in hours_per_unit.items() if hours > 0]
                line_value = str(row[header_columns["linea produccion"]] or "").strip()

                scheduled_values = [
                    (column_index, week_number, row[column_index] if column_index < len(row) else None)
                    for column_index, week_number in week_columns
                    if column_index < len(row) and row[column_index] is not None and str(row[column_index]).strip()
                ]
                if not scheduled_values:
                    try:
                        quantity = _scheduled_quantity(row[header_columns["restante"]])
                        delivery_date = row[header_columns["fecha"]]
                        if not isinstance(delivery_date, date):
                            raise ValueError(f"Fecha de entrega no válida: {delivery_date}")
                        iso_date = delivery_date.isocalendar()
                        delivery_week = f"{iso_date.year}-W{iso_date.week:02d}"
                        safe_of = re.sub(r"[^A-Za-z0-9_.-]+", "_", source_of).strip("_.-") or "OF"
                        code = f"{safe_of}-R{row_number}-{delivery_week.replace('-', '')}"
                        orders.append(
                            {
                                "code": code,
                                "source_of": source_of,
                                "client": client,
                                "project": project,
                                "quantity": quantity,
                                "delivery_week": delivery_week,
                                "lead_time_weeks": 2,
                                "required_skill": "general",
                                "hours_per_unit": hours_per_unit,
                                "active_phases": active_phases,
                                "production_line": line_value,
                                "row": row_number,
                            }
                        )
                    except (TypeError, ValueError) as error:
                        errors.append({"row": row_number, "error": str(error)})
                    continue

                for column_index, week_number, value in scheduled_values:
                    delivery_year = 2025 if week_number == 52 and column_index < first_week_one_column else 2026
                    delivery_week = f"{delivery_year}-W{week_number:02d}"
                    try:
                        quantity = _scheduled_quantity(value)
                    except ValueError as error:
                        errors.append({"row": row_number, "week": delivery_week, "error": str(error)})
                        continue

                    safe_of = re.sub(r"[^A-Za-z0-9_.-]+", "_", source_of).strip("_.-") or "OF"
                    code = f"{safe_of}-R{row_number}-{delivery_year}W{week_number:02d}"
                    orders.append(
                        {
                            "code": code,
                            "source_of": source_of,
                            "client": client,
                            "project": project,
                            "quantity": quantity,
                            "delivery_week": delivery_week,
                            "lead_time_weeks": 2,
                            "required_skill": "general",
                            "hours_per_unit": hours_per_unit,
                            "active_phases": active_phases,
                            "production_line": line_value,
                            "row": row_number,
                        }
                    )
            return orders, errors
        return None, []
    finally:
        workbook.close()